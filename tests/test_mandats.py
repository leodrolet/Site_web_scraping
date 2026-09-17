"""Tests ASGI sans service externe, dans une base SQLite temporaire.

Lancer : .venv/bin/python -m unittest discover -s tests -v
"""
import asyncio
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from datetime import datetime, timedelta
from unittest.mock import Mock, patch
from urllib.parse import urlencode

# L'app charge .env au démarrage : ces valeurs explicites empêchent tout accès
# à la base et aux fournisseurs du propriétaire, même si .env est présent.
RACINE = Path(__file__).resolve().parents[1]
TEMP = tempfile.TemporaryDirectory(prefix="prospect-mandats-tests-")
for cle in ("DATABASE_URL", "VERCEL", "HUNTER_API_KEY", "APOLLO_API_KEY", "SERPAPI_KEY", "RESEND_API_KEY"):
    os.environ[cle] = ""
os.environ.update(SECRET_KEY="cle-uniquement-pour-tests", COOKIE_SECURE="false", EXIGER_CONFIRMATION_EMAIL="false")
os.chdir(TEMP.name)
import sys
sys.path.insert(0, str(RACINE))

from main import app
import database
from database import Base, SessionLocal, Utilisateur, DossierRecherche, EntrepriseDossier, ImportMandat, HistoriqueRecherche
from auth import _serializer
import mandats
import mandats_recherche
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, inspect, text


async def requete(method, path, body, content_type, cookie):
    sortie = []
    envoye = False
    chemin, _, query = path.partition("?")
    async def receive():
        nonlocal envoye
        if not envoye:
            envoye = True
            return {"type": "http.request", "body": body, "more_body": False}
        await asyncio.Event().wait()
    async def send(message):
        sortie.append(message)
    scope = dict(type="http", asgi={"version": "3.0", "spec_version": "2.4"}, http_version="1.1",
                 method=method, scheme="http", path=chemin, raw_path=chemin.encode(), query_string=query.encode(),
                 root_path="", headers=[(b"host", b"testserver"), (b"content-type", content_type.encode()),
                                        (b"cookie", cookie.encode())], client=("127.0.0.1", 1), server=("testserver", 80))
    await asyncio.wait_for(app(scope, receive, send), 10)
    entetes = next(m for m in sortie if m["type"] == "http.response.start")
    return (entetes["status"], b"".join(m.get("body", b"") for m in sortie if m["type"] == "http.response.body"),
            dict(entetes["headers"]))


class MandatsTest(unittest.TestCase):
    def setUp(self):
        with database.engine.begin() as conn:
            # Refuse de nettoyer une base réutilisée par un autre environnement.
            chemins = conn.execute(text("PRAGMA database_list")).fetchall()
            if not all(not row[2] or Path(row[2]).resolve().is_relative_to(Path(TEMP.name).resolve()) for row in chemins):
                raise RuntimeError("Les tests doivent utiliser exclusivement leur base temporaire.")
            for table in reversed(Base.metadata.sorted_tables):
                conn.execute(table.delete())
        self.db = SessionLocal()
        self.admin = Utilisateur(email="admin@example.org", mot_de_passe_hash="test", admin=True)
        self.autre = Utilisateur(email="autre@example.org", mot_de_passe_hash="test", admin=True)
        self.public = Utilisateur(email="public@example.org", mot_de_passe_hash="test", admin=False)
        self.db.add_all([self.admin, self.autre, self.public])
        self.db.commit()
        self.d = DossierRecherche(utilisateur_id=self.admin.id, nom="Mandat Lyse")
        self.db.add(self.d)
        self.db.commit()
        self.base = f"/admin/mandats/{self.d.id}"

    def tearDown(self):
        self.db.close()

    def call(self, method, path, data=None, user="admin", fichier=None, csrf=True):
        u = {"admin": self.admin, "autre": self.autre, "public": self.public, "anonyme": None}[user]
        cookie = f"session={_serializer.dumps(u.id)}; csrf_token=jeton-test" if u else ""
        champs = dict(data or {})
        if method == "POST" and csrf:
            champs["csrf_token"] = "jeton-test"
        if fichier:
            boundary = "TEST-BOUNDARY"
            parts = [(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n').encode() for k, v in champs.items()]
            nom, contenu = fichier
            parts.append((f'--{boundary}\r\nContent-Disposition: form-data; name="fichier"; filename="{nom}"\r\nContent-Type: application/octet-stream\r\n\r\n').encode() + contenu + b"\r\n")
            body = b"".join(parts) + f"--{boundary}--\r\n".encode()
            ctype = "multipart/form-data; boundary=" + boundary
        else:
            body, ctype = urlencode(champs).encode(), "application/x-www-form-urlencoded"
        return asyncio.run(requete(method, path, body, ctype, cookie))

    def org(self, nom="Atelier Test", **champs):
        e = EntrepriseDossier(dossier_id=self.d.id, nom=nom, cle=mandats.normaliser(nom), **champs)
        self.db.add(e)
        self.db.commit()
        return e

    def fiche(self, e):
        return self.base + f"/organisations/{e.id}"

    def importer(self, contenu=None):
        contenu = contenu or "Raison sociale;Catégorie;Pays;Prix;Mémo;Autre\nAtelier ÉCO;Architecture;Canada;Excellence;Carbone zéro;Valeur conservée\n ATELIER ECO ;Architecture;Canada;Prix durable;Bureaux;Seconde valeur\nUniversité Test;Enseignement;France;;;\n"
        code, _, h = self.call("POST", self.base + "/preparer", fichier=("lyse.csv", contenu.encode("utf-8-sig")))
        self.assertEqual(code, 303)
        chemin = h[b"location"].decode()
        mapping = {"nom": "0", "secteur": "1", "pays": "2", "prix": "3", "note": "4"}
        code, _, h = self.call("POST", chemin + "/mapper", mapping)
        self.assertEqual(code, 303)
        imp = self.db.get(ImportMandat, chemin.rsplit("/", 1)[1])
        self.db.refresh(imp)
        return chemin, imp, json.loads(imp.mapping_json)["jeton"]

    def test_routes_admin_et_proprietaire(self):
        e = self.org()
        routes = [("GET", "/admin/mandats"), ("POST", "/admin/mandats"),
                  ("GET", self.base), ("POST", self.base + "/parametres"),
                  ("POST", self.base + "/organisations"), ("POST", self.base + "/preparer"),
                  ("GET", self.fiche(e)), ("POST", self.fiche(e)),
                  ("POST", self.fiche(e) + "/contacts"), ("POST", self.fiche(e) + "/rechercher"),
                  ("GET", self.base + "/export"), ("GET", "/app/dossiers"),
                  ("POST", "/app/dossiers/1/importer"),
                  ("GET", self.base + "/imports/abc"), ("POST", self.base + "/imports/abc/mapper"),
                  ("GET", self.base + "/imports/abc/apercu"), ("POST", self.base + "/imports/abc/confirmer")]
        for method, chemin in routes:
            for user, destination in (("anonyme", b"/login"), ("public", b"/app")):
                with self.subTest(path=chemin, user=user, method=method):
                    code, _, h = self.call(method, chemin, {"nom": "Test"}, user=user)
                    self.assertEqual(code, 303)
                    self.assertEqual(h[b"location"], destination)
        for method, chemin in [("GET", self.base), ("GET", self.fiche(e)), ("GET", self.base + "/export"),
                               ("POST", self.fiche(e) + "/contacts")]:
            self.assertEqual(self.call(method, chemin, user="autre")[0], 404)
        self.assertNotIn(b'/admin/mandats', self.call("GET", "/app", user="public")[1])
        self.assertEqual(self.call("GET", "/app/dossiers", user="admin")[2][b"location"], b"/admin/mandats")

    def test_csrf_lecture_et_sauvegarde(self):
        self.assertEqual(self.call("POST", "/admin/mandats", {"nom": "Non"}, csrf=False)[0], 403)
        e = self.org()
        for suffix in ("", "/contacts", "/rechercher"):
            self.assertEqual(self.call("POST", self.fiche(e) + suffix, csrf=False)[0], 403)
        code, corps, _ = self.call("GET", self.fiche(e))
        self.assertEqual(code, 200)
        self.assertIn("Ajouter un contact".encode(), corps)
        donnees = dict(nom="Atelier corrigé", secteur="Design", pays="Canada", site="exemple.org", prix="Prix 2026", statut="a_verifier", revision=e.revision)
        self.assertEqual(self.call("POST", self.fiche(e), donnees)[0], 303)
        self.db.refresh(e)
        self.assertEqual(e.site, "https://exemple.org")
        self.assertEqual(e.prix, "Prix 2026")
        self.assertEqual(self.call("POST", self.fiche(e), donnees)[0], 409)
        code, corps, _ = self.call("POST", self.fiche(e), dict(donnees, revision=e.revision, site="javascript:alert(1)"))
        self.assertEqual(code, 400)
        self.db.refresh(e)
        self.assertEqual(e.site, "https://exemple.org")

    def test_import_apercu_doublons_et_conservation(self):
        chemin, imp, jeton = self.importer()
        self.assertEqual(self.db.query(EntrepriseDossier).count(), 0)
        for suffix in ("", "/apercu"):
            code, corps, _ = self.call("GET", chemin + suffix)
            self.assertEqual(code, 200)
        code, corps, _ = self.call("POST", chemin + "/confirmer", dict(jeton=jeton, doublons="completer"))
        self.assertEqual(code, 200, corps[:400])
        es = self.db.query(EntrepriseDossier).all()
        self.assertEqual(len(es), 2)
        eco = es[0]
        self.assertEqual(eco.nom, "Atelier ÉCO")
        self.assertIn("Prix durable", eco.prix)
        self.assertIn("Bureaux", eco.note)
        self.assertIn("Valeur conservée", eco.donnees_import_json)
        self.assertEqual(len(json.loads(eco.donnees_import_json)), 2)
        self.assertEqual(self.call("POST", chemin + "/confirmer", dict(jeton=jeton))[0], 410)
        self.assertEqual(self.call("GET", self.fiche(eco))[0], 200)

    def test_import_invalide_expire_et_xlsx(self):
        code, corps, _ = self.call("POST", self.base + "/preparer", fichier=("ancien.xls", b"bad"))
        self.assertEqual(code, 400)
        self.assertIn("manuellement".encode(), corps)
        w = Workbook(); w.active.title = "Info"; s = w.create_sheet("Liste")
        s.append(["Titre du document"]); s.append(["company", "country", "website"]); s.append(["Engineering Test", "UK", "test.example"])
        b = io.BytesIO(); w.save(b)
        code, _, h = self.call("POST", self.base + "/preparer", dict(feuille=2, ligne_entete=2), fichier=("lyse.xlsx", b.getvalue()))
        self.assertEqual(code, 303)
        chemin = h[b"location"].decode()
        self.assertEqual(self.call("GET", chemin)[0], 200)
        imp = self.db.get(ImportMandat, chemin.rsplit("/", 1)[1])
        imp.cree_le = datetime.utcnow() - timedelta(days=2); self.db.commit()
        self.assertEqual(self.call("GET", chemin)[0], 410)
        s.append(["=1+1"]); b = io.BytesIO(); w.save(b)
        self.assertEqual(self.call("POST", self.base + "/preparer", dict(feuille=2, ligne_entete=2), fichier=("formules.xlsx", b.getvalue()))[0], 400)

    def test_cinq_contacts_selection_validation_export(self):
        e = self.org(nom="=Organisation", statut="termine", secteur="Construction", pays="Canada")
        self.org(nom="Sans contact", statut="termine")
        for i in range(5):
            self.db.refresh(e)
            f = dict(nom=f"Contact {i}", poste="Directeur marketing", region="Canada", langue="Anglais (profil)",
                     source="https://example.org/team", verification="verifie", date_verification="2026-01-01",
                     courriel_statut="non_trouve", revision=e.revision, note="=1+1")
            if i < 3:
                f["retenu"] = "on"
            self.assertEqual(self.call("POST", self.fiche(e) + "/contacts", f)[0], 303)
        self.db.refresh(e)
        cs = mandats.contacts_de(e)
        self.assertEqual(len(cs), 5)
        mauvais = dict(nom="Contact 0", contact_id=cs[0]["id"], courriel="fake@example.org", courriel_statut="trouve", revision=e.revision)
        self.assertEqual(self.call("POST", self.fiche(e) + "/contacts", mauvais)[0], 400)
        self.db.refresh(e)
        self.assertFalse(mandats.contacts_de(e)[0]["courriel"])
        self.assertEqual(self.call("GET", self.fiche(e))[0], 200)
        code, contenu, _ = self.call("GET", self.base + "/export")
        self.assertEqual(code, 200)
        wb = load_workbook(io.BytesIO(contenu))
        self.assertEqual(wb.sheetnames, ["Contacts retenus", "Organisations", "Bilan"])
        self.assertEqual(wb["Contacts retenus"].max_row, 4)
        self.assertEqual(wb["Contacts retenus"]["A2"].data_type, "s")
        self.assertEqual(wb["Contacts retenus"]["P2"].data_type, "s")
        self.assertEqual(wb["Organisations"]["M3"].value, "Aucun contact pertinent retenu")

    def test_api_bornee_et_suggestions_non_verifiees(self):
        e = self.org(site="https://example.org")
        rep = Mock(status_code=200)
        rep.json.return_value = {"data": {"domain": "example.org", "emails": [
            {"first_name": "Personne", "last_name": "Test", "position": "Marketing Director", "department": "marketing",
             "value": "personne@example.org", "sources": [{"uri": "https://example.org/team"}]}] * 10}}
        with patch.dict(os.environ, {"HUNTER_API_KEY": "secret-fictif"}), patch("recherche.requests.get", return_value=rep) as get:
            code, corps, _ = self.call("POST", self.fiche(e) + "/rechercher", dict(service="hunter", departement="Marketing", region="Toutes", revision=e.revision))
            self.assertEqual(code, 200)
            self.assertEqual(get.call_count, 1)
            self.assertNotIn(b"secret-fictif", corps)
        self.db.refresh(e)
        cs = mandats.contacts_de(e)
        self.assertEqual(len(cs), 1)
        self.assertEqual(cs[0]["verification"], "a_verifier")
        self.assertEqual(cs[0]["courriel_statut"], "a_verifier")
        self.assertEqual(cs[0]["date_verification"], "")
        self.assertFalse(cs[0]["retenu"])
        rep.json.return_value = {"organic_results": [{"link": "https://www.linkedin.com/in/piste-test"}]}
        with patch.dict(os.environ, {"SERPAPI_KEY": "secret-fictif"}), patch("recherche.requests.get", return_value=rep) as get:
            self.assertEqual(self.call("POST", self.fiche(e) + "/rechercher", dict(service="serpapi", departement="Partenariats", region="Toutes", revision=e.revision))[0], 200)
            self.assertEqual(get.call_count, 1)
        self.db.refresh(e)
        self.assertEqual(len(mandats.contacts_de(e)), 1)
        self.assertEqual(len(json.loads(e.pistes_json)), 1)
        with patch.dict(os.environ, {"HUNTER_API_KEY": "secret-fictif"}), patch("recherche.requests.get", side_effect=RuntimeError("secret-fictif")):
            code, corps, _ = self.call("POST", self.fiche(e) + "/rechercher", dict(service="hunter", departement="Marketing", region="Toutes", revision=e.revision))
            self.assertEqual(code, 400)
            self.assertNotIn(b"secret-fictif", corps)

    def test_sans_api_et_export_donnees_compte(self):
        e = self.org(pays="Canada", prix="Excellence", durabilite="Carbone zéro")
        self.assertEqual(self.call("POST", self.fiche(e) + "/rechercher", dict(service="hunter", revision=0))[0], 400)
        self.assertEqual(self.call("GET", self.fiche(e))[0], 200)
        self.assertEqual(self.call("POST", self.base + "/parametres", dict(nom="Nouveau nom", temps_minutes=75, budget_minutes=420))[0], 303)
        data = json.loads(self.call("GET", "/mes-donnees")[1])
        self.assertEqual(data["dossiers"][0]["temps_minutes"], 75)
        self.assertEqual(data["dossiers"][0]["entreprises"][0]["prix"], "Excellence")

    def test_creation_saisie_manuelle_et_reprise(self):
        code, _, h = self.call("POST", "/admin/mandats", {"nom": "  Nouveau mandat  "})
        self.assertEqual(code, 303)
        cible = h[b"location"].decode()
        code, _, h = self.call("POST", cible + "/organisations", dict(
            nom="Organisation manuelle", secteur="Génie-conseil", pays="Canada",
            taille="grande", sources="https://example.org/about", durabilite="Carbone zéro"))
        self.assertEqual(code, 303)
        fiche = h[b"location"].decode()
        self.assertEqual(self.call("GET", fiche)[0], 200)
        self.assertEqual(self.call("GET", cible)[0], 200)
        self.assertEqual(self.call("POST", cible + "/organisations", dict(nom="organisation MANUELLE", pays="Canada"))[0], 400)

    def test_csv_separateurs_encodages_et_mapping_invalide(self):
        from types import SimpleNamespace
        for sep in (",", ";", "\t"):
            for enc in ("utf-8-sig", "cp1252"):
                with self.subTest(separateur=sep, encodage=enc):
                    contenu = sep.join(["Entreprise", "Pays"]) + "\n" + sep.join(["École Éco", "Canada"])
                    fichier = SimpleNamespace(filename="liste.csv", file=io.BytesIO(contenu.encode(enc)))
                    self.assertEqual(mandats.lire_fichier(fichier)["lignes"][0][0], "École Éco")
        chemin, imp, _ = self.importer()
        self.assertEqual(self.call("POST", chemin + "/mapper", dict(nom="0", pays="0"))[0], 400)
        self.assertEqual(self.call("POST", chemin + "/mapper", dict(nom="99"))[0], 400)
        self.assertEqual(self.call("GET", chemin, user="autre")[0], 404)
        self.assertEqual(self.call("POST", chemin + "/confirmer", csrf=False)[0], 403)

    def test_import_fusion_existante_et_exclusion_invalide(self):
        e = self.org(nom="Atelier ÉCO", pays="Canada", sources="https://example.org/a",
                     contacts_json='[{"nom":"Contact existant"}]', statut="termine")
        contenu = "Raison sociale;Catégorie;Pays;Prix;Mémo;Autre\nATELIER ECO;Design;Canada;Prix;Note;Original\n;Architecture;Canada;;;\n"
        chemin, _, jeton = self.importer(contenu)
        self.assertEqual(self.call("POST", chemin + "/confirmer", dict(jeton=jeton, doublons="completer"))[0], 400)
        self.assertEqual(self.call("POST", chemin + "/confirmer", dict(jeton=jeton, doublons="completer", ignorer_invalides="on"))[0], 200)
        self.db.refresh(e)
        self.assertEqual(e.prix, "Prix")
        self.assertEqual(e.statut, "termine")
        self.assertEqual(mandats.contacts_de(e)[0]["nom"], "Contact existant")
        self.assertFalse(mandats.contacts_de(e)[0]["retenu"])
        self.assertEqual(self.db.query(EntrepriseDossier).count(), 1)

    def test_modifier_contact_et_validation(self):
        e = self.org()
        chemin = self.fiche(e) + "/contacts"
        self.assertEqual(self.call("POST", chemin, dict(nom="Personne Test", revision=0))[0], 303)
        self.db.refresh(e)
        c = mandats.contacts_de(e)[0]
        original = dict(contact_id=c["id"], nom="Personne Test", revision=e.revision,
                        poste="Communications", source="https://example.org/equipe",
                        date_verification="2026-01-01", verification="verifie",
                        courriel="contact@example.org", courriel_statut="trouve",
                        source_courriel="https://example.org/contact", retenu="on")
        for modifications in (dict(verification="ecarte"), dict(date_verification="2999-01-01"),
                              dict(source=""), dict(courriel="invalide"), dict(courriel_statut="non_trouve")):
            self.assertEqual(self.call("POST", chemin, dict(original, **modifications))[0], 400)
        self.assertEqual(self.call("POST", chemin, original)[0], 303)
        self.db.refresh(e)
        c = mandats.contacts_de(e)[0]
        self.assertEqual(c["verification"], "verifie")
        self.assertEqual(c["poste"], "Communications")
        self.assertEqual(c["source"], "https://example.org/equipe")
        self.assertTrue(c["retenu"])
        original.update(revision=e.revision, verification="ecarte")
        original.pop("retenu")
        self.assertEqual(self.call("POST", chemin, original)[0], 303)
        self.db.refresh(e)
        self.assertFalse(mandats.contacts_de(e)[0]["retenu"])

    def test_apollo_une_page_et_aucun_courriel_genere(self):
        e = self.org()
        rep = Mock(status_code=200)
        rep.json.return_value = {"people": [dict(first_name="Personne", last_name="Test", title="Partnerships Director",
            linkedin_url="https://www.linkedin.com/in/exemple", email="email_not_unlocked", city="Montréal",
            organization={"name": "Organisation exemple"})] * 100, "pagination": {"total_pages": 25}}
        with patch.dict(os.environ, {"APOLLO_API_KEY": "secret-test"}), patch("recherche.requests.post", return_value=rep) as post:
            resultats, pistes, _ = mandats_recherche.chercher(e, "apollo", "Partenariats", "Canada")
        self.assertEqual(post.call_count, 1)
        self.assertEqual(post.call_args.kwargs["json"]["page"], 1)
        self.assertFalse(pistes)
        self.assertTrue(resultats)
        self.assertTrue(all(c["courriel"] == "" and c["courriel_statut"] == "non_trouve" for c in resultats))
        self.assertTrue(all(c["verification"] == "a_verifier" and not c["retenu"] for c in resultats))

    def test_suppression_cascade_sans_autres_comptes(self):
        self.org(contacts_json='[{"nom":"Contact privé"}]')
        self.importer()
        self.db.add(HistoriqueRecherche(utilisateur_id=self.admin.id, entreprise="Historique"))
        self.db.commit()
        self.db.delete(self.admin)
        self.db.commit()
        self.assertEqual(self.db.query(EntrepriseDossier).count(), 0)
        self.assertEqual(self.db.query(ImportMandat).count(), 0)
        self.assertEqual(self.db.query(DossierRecherche).count(), 0)
        self.assertEqual(self.db.query(HistoriqueRecherche).count(), 0)
        self.assertEqual(self.db.query(Utilisateur).count(), 2)

    def test_migration_additive_idempotente(self):
        ancien = create_engine("sqlite:///:memory:")
        Base.metadata.tables["utilisateurs"].create(ancien)
        Base.metadata.tables["historique_recherches"].create(ancien)
        with ancien.begin() as c:
            c.execute(text("INSERT INTO utilisateurs (id,email,mot_de_passe_hash,admin) VALUES (7,'ancien@example.org','hash-conserve',1)"))
            c.execute(text("INSERT INTO historique_recherches (id,utilisateur_id,entreprise) VALUES (9,7,'Historique conservé')"))
            c.execute(text("CREATE TABLE dossiers_recherche (id INTEGER PRIMARY KEY, utilisateur_id INTEGER NOT NULL, nom VARCHAR(160), debut DATETIME, fin DATETIME)"))
            c.execute(text("CREATE TABLE entreprises_dossier (id INTEGER PRIMARY KEY, dossier_id INTEGER, nom VARCHAR(255), cle VARCHAR(255), secteur VARCHAR(160), region VARCHAR(160), site VARCHAR(500), statut VARCHAR(20), note TEXT, contacts_json TEXT, mise_a_jour DATETIME)"))
            c.execute(text("INSERT INTO dossiers_recherche (id,utilisateur_id,nom) VALUES (3,7,'Ancien dossier')"))
            c.execute(text("INSERT INTO entreprises_dossier (id,dossier_id,nom,contacts_json) VALUES (4,3,'Ancienne organisation',:contacts)"), {"contacts": '[{"nom":"Contact conservé"}]'})
        with patch.object(database, "engine", ancien):
            database.init_db(); database.init_db()
        with ancien.connect() as c:
            self.assertEqual(c.execute(text("SELECT mot_de_passe_hash FROM utilisateurs WHERE id=7")).scalar(), "hash-conserve")
            self.assertEqual(c.execute(text("SELECT entreprise FROM historique_recherches WHERE id=9")).scalar(), "Historique conservé")
            self.assertIn("Contact conservé", c.execute(text("SELECT contacts_json FROM entreprises_dossier WHERE id=4")).scalar())
            self.assertEqual(c.execute(text("SELECT revision FROM entreprises_dossier WHERE id=4")).scalar(), 0)
        self.assertIn("imports_mandat", inspect(ancien).get_table_names())
        ancien.dispose()


if __name__ == "__main__":
    unittest.main()
