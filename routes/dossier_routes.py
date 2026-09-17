"""Dossiers de recherche manuelle, séparés des recherches API et de leur quota."""

import csv
import io
import json
import re
import unicodedata
from datetime import datetime
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import RedirectResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from auth import exiger_connexion, valider_csrf
from database import DossierRecherche, EntrepriseDossier, Utilisateur, get_db
from templating import rendre

router = APIRouter(prefix="/app/dossiers")
MAX_FICHIER = 2_000_000
MAX_LIGNES = 2000


def _cle(nom):
    texte = unicodedata.normalize("NFKD", nom.casefold())
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", texte)


def _url(valeur):
    valeur = valeur.strip()
    if not valeur:
        return ""
    p = urlparse(valeur)
    if p.scheme not in ("http", "https") or not p.netloc or p.username or p.password:
        raise ValueError("Utilisez un lien complet commençant par https:// ou http://.")
    return valeur[:500]


def _cellule(valeur):
    """Empêche les valeurs importées de devenir des formules Excel."""
    texte = str(valeur or "")
    return "'" + texte if texte.lstrip().startswith(("=", "+", "-", "@")) else texte


def _lire(fichier):
    nom = (fichier.filename or "").lower()
    if not nom.endswith((".csv", ".xlsx")):
        raise ValueError("Importez un fichier CSV ou Excel (.xlsx).")
    brut = fichier.file.read(MAX_FICHIER + 1)
    if len(brut) > MAX_FICHIER:
        raise ValueError("Le fichier dépasse 2 Mo.")
    if nom.endswith(".xlsx"):
        try:
            livre = load_workbook(io.BytesIO(brut), read_only=True, data_only=True)
            feuille = livre.active
            lignes = feuille.iter_rows(values_only=True)
            entetes = next(lignes, None)
            if not entetes:
                raise ValueError("Le fichier est vide.")
            donnees = []
            for n, ligne in enumerate(lignes, 1):
                if n > MAX_LIGNES:
                    raise ValueError("Limite de 2 000 entreprises par import.")
                donnees.append(ligne)
            livre.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Fichier Excel illisible.") from exc
    else:
        try:
            contenu = brut.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Le CSV doit être encodé en UTF-8.") from exc
        echantillon = contenu[:2048]
        separateur = ";" if echantillon.count(";") > echantillon.count(",") else ","
        lignes = csv.reader(io.StringIO(contenu), delimiter=separateur)
        entetes = next(lignes, None)
        if not entetes:
            raise ValueError("Le fichier est vide.")
        donnees = []
        for n, ligne in enumerate(lignes, 1):
            if n > MAX_LIGNES:
                raise ValueError("Limite de 2 000 entreprises par import.")
            donnees.append(ligne)

    noms = [_cle(str(h or "")) for h in entetes]
    def colonne(*possibles):
        return next((i for i, h in enumerate(noms) if h in possibles), None)
    idx_nom = colonne("entreprise", "nomentreprise", "nomdelentreprise", "company", "companyname", "organisation", "organization", "nom")
    if idx_nom is None:
        raise ValueError("Ajoutez une colonne « entreprise » ou « company » au fichier.")
    indices = [idx_nom, colonne("secteur", "industrie", "industry", "sector"),
               colonne("region", "ville", "city", "province", "location"),
               colonne("site", "siteweb", "website", "url", "domaine")]
    resultat = []
    for ligne in donnees:
        champs = [str(ligne[i] or "").strip() if i is not None and i < len(ligne) else ""
                  for i in indices]
        if champs[0]:
            resultat.append(champs)
    return resultat


def _dossier(db, utilisateur, dossier_id):
    return (db.query(DossierRecherche)
            .filter(DossierRecherche.id == dossier_id,
                    DossierRecherche.utilisateur_id == utilisateur.id).first())


def _page(request, db, utilisateur, dossier=None, erreur=None, message=None):
    dossiers = (db.query(DossierRecherche)
                .filter(DossierRecherche.utilisateur_id == utilisateur.id)
                .order_by(DossierRecherche.debut.desc()).all())
    entreprises = []
    traitees = 0
    if dossier:
        entreprises = (db.query(EntrepriseDossier)
                       .filter(EntrepriseDossier.dossier_id == dossier.id)
                       .order_by(EntrepriseDossier.id).all())
        traitees = sum(e.statut == "termine" for e in entreprises)
    return rendre(request, "dossiers.html", utilisateur=utilisateur,
                  dossiers=dossiers, dossier=dossier, entreprises=entreprises,
                  traitees=traitees, erreur=erreur, message=message)


@router.get("")
def liste(request: Request, utilisateur: Utilisateur = Depends(exiger_connexion),
          db: Session = Depends(get_db)):
    return _page(request, db, utilisateur)


@router.post("")
def creer(request: Request, nom: str = Form(...), csrf_token: str = Form(""),
          utilisateur: Utilisateur = Depends(exiger_connexion),
          db: Session = Depends(get_db)):
    if not valider_csrf(request, csrf_token):
        return _page(request, db, utilisateur, erreur="Session expirée. Réessayez.")
    nom = nom.strip()[:160]
    if not nom:
        return _page(request, db, utilisateur, erreur="Donnez un nom au dossier.")
    dossier = DossierRecherche(utilisateur_id=utilisateur.id, nom=nom)
    db.add(dossier)
    db.commit()
    return RedirectResponse(f"/app/dossiers/{dossier.id}", status_code=303)


@router.get("/{dossier_id}")
def detail(dossier_id: int, request: Request,
           utilisateur: Utilisateur = Depends(exiger_connexion),
           db: Session = Depends(get_db)):
    dossier = _dossier(db, utilisateur, dossier_id)
    if not dossier:
        return RedirectResponse("/app/dossiers", status_code=303)
    return _page(request, db, utilisateur, dossier)


@router.post("/{dossier_id}/importer")
def importer(dossier_id: int, request: Request, fichier: UploadFile = File(...),
             csrf_token: str = Form(""),
             utilisateur: Utilisateur = Depends(exiger_connexion),
             db: Session = Depends(get_db)):
    dossier = _dossier(db, utilisateur, dossier_id)
    if not dossier:
        return RedirectResponse("/app/dossiers", status_code=303)
    if not valider_csrf(request, csrf_token):
        return _page(request, db, utilisateur, dossier, erreur="Session expirée. Réessayez.")
    try:
        lignes = _lire(fichier)
    except ValueError as exc:
        return _page(request, db, utilisateur, dossier, erreur=str(exc))
    connues = {e.cle for e in dossier.entreprises}
    ajout = 0
    for nom, secteur, region, site in lignes:
        cle = _cle(nom)
        if not cle or cle in connues:
            continue
        try:
            site = _url(site) if site.startswith(("http://", "https://")) else ""
        except ValueError:
            site = ""
        db.add(EntrepriseDossier(dossier_id=dossier.id, nom=nom[:255], cle=cle[:255],
                                 secteur=secteur[:160], region=region[:160], site=site))
        connues.add(cle)
        ajout += 1
    db.commit()
    return _page(request, db, utilisateur, dossier,
                 message=f"{ajout} entreprise(s) ajoutée(s). {len(lignes) - ajout} doublon(s) ou ligne(s) ignorée(s).")


@router.get("/{dossier_id}/entreprises/{entreprise_id}")
def fiche(dossier_id: int, entreprise_id: int, request: Request,
          utilisateur: Utilisateur = Depends(exiger_connexion),
          db: Session = Depends(get_db)):
    dossier = _dossier(db, utilisateur, dossier_id)
    if not dossier:
        return RedirectResponse("/app/dossiers", status_code=303)
    entreprise = (db.query(EntrepriseDossier)
                  .filter(EntrepriseDossier.id == entreprise_id,
                          EntrepriseDossier.dossier_id == dossier.id).first())
    if not entreprise:
        return RedirectResponse(f"/app/dossiers/{dossier.id}", status_code=303)
    return rendre(request, "dossier_fiche.html", utilisateur=utilisateur,
                  dossier=dossier, entreprise=entreprise,
                  contacts=(json.loads(entreprise.contacts_json or "[]") + [{}, {}])[:2])


@router.post("/{dossier_id}/entreprises/{entreprise_id}")
async def sauvegarder(dossier_id: int, entreprise_id: int, request: Request,
                      utilisateur: Utilisateur = Depends(exiger_connexion),
                      db: Session = Depends(get_db)):
    dossier = _dossier(db, utilisateur, dossier_id)
    entreprise = (db.query(EntrepriseDossier)
                  .filter(EntrepriseDossier.id == entreprise_id,
                          EntrepriseDossier.dossier_id == dossier_id).first()) if dossier else None
    if not entreprise:
        return RedirectResponse("/app/dossiers", status_code=303)
    form = await request.form()
    saisie_contacts = [dict(nom=str(form.get(f"nom_{i}", "")),
                           poste=str(form.get(f"poste_{i}", "")),
                           source=str(form.get(f"source_{i}", "")),
                           courriel=str(form.get(f"courriel_{i}", "")),
                           source_courriel=str(form.get(f"source_courriel_{i}", "")))
                       for i in (1, 2)]
    if not valider_csrf(request, str(form.get("csrf_token", ""))):
        erreur = "Session expirée. Réessayez."
    else:
        try:
            entreprise.secteur = str(form.get("secteur", "")).strip()[:160]
            entreprise.region = str(form.get("region", "")).strip()[:160]
            entreprise.site = _url(str(form.get("site", "")))
            entreprise.note = str(form.get("note", "")).strip()[:2000]
            contacts = []
            for i in (1, 2):
                nom = str(form.get(f"nom_{i}", "")).strip()[:160]
                poste = str(form.get(f"poste_{i}", "")).strip()[:160]
                source = _url(str(form.get(f"source_{i}", "")))
                courriel = str(form.get(f"courriel_{i}", "")).strip()[:254]
                source_courriel = _url(str(form.get(f"source_courriel_{i}", "")))
                if (nom or poste or courriel) and not source:
                    raise ValueError(f"Contact {i} : ajoutez un lien qui confirme le poste actuel.")
                if courriel and (not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", courriel) or not source_courriel):
                    raise ValueError(f"Contact {i} : indiquez un courriel valide et sa source fiable.")
                contacts.append(dict(nom=nom, poste=poste, source=source,
                                     courriel=courriel, source_courriel=source_courriel))
            statut = str(form.get("statut", "en_cours"))
            if statut not in ("a_faire", "en_cours", "termine"):
                raise ValueError("Statut invalide.")
            entreprise.contacts_json = json.dumps(contacts, ensure_ascii=False)
            entreprise.statut = statut
            entreprise.mise_a_jour = datetime.utcnow()
            db.commit()
            return RedirectResponse(f"/app/dossiers/{dossier.id}", status_code=303)
        except ValueError as exc:
            db.rollback()
            erreur = str(exc)
    # Représente la saisie même si une validation a échoué.
    entreprise.secteur = str(form.get("secteur", ""))
    entreprise.region = str(form.get("region", ""))
    entreprise.site = str(form.get("site", ""))
    entreprise.note = str(form.get("note", ""))
    entreprise.statut = str(form.get("statut", "en_cours"))
    return rendre(request, "dossier_fiche.html", utilisateur=utilisateur,
                  dossier=dossier, entreprise=entreprise,
                  contacts=saisie_contacts,
                  erreur=erreur)


@router.post("/{dossier_id}/terminer")
def terminer(dossier_id: int, request: Request, csrf_token: str = Form(""),
             utilisateur: Utilisateur = Depends(exiger_connexion),
             db: Session = Depends(get_db)):
    dossier = _dossier(db, utilisateur, dossier_id)
    if not dossier:
        return RedirectResponse("/app/dossiers", status_code=303)
    if not valider_csrf(request, csrf_token):
        return _page(request, db, utilisateur, dossier, erreur="Session expirée. Réessayez.")
    if dossier.fin is None:
        dossier.fin = datetime.utcnow()
        db.commit()
    return RedirectResponse(f"/app/dossiers/{dossier.id}", status_code=303)


@router.get("/{dossier_id}/export")
def exporter(dossier_id: int, utilisateur: Utilisateur = Depends(exiger_connexion),
             db: Session = Depends(get_db)):
    dossier = _dossier(db, utilisateur, dossier_id)
    if not dossier:
        return RedirectResponse("/app/dossiers", status_code=303)
    entreprises = (db.query(EntrepriseDossier)
                   .filter(EntrepriseDossier.dossier_id == dossier.id,
                           EntrepriseDossier.statut == "termine")
                   .order_by(EntrepriseDossier.id).all())
    livre = Workbook()
    feuille = livre.active
    feuille.title = "Contacts vérifiés"
    colonnes = ["Entreprise", "Secteur", "Nom", "Poste", "Région", "Courriel disponible",
                "Lien de la source", "Source du courriel", "Note"]
    feuille.append(colonnes)
    for entreprise in entreprises:
        contacts = [c for c in json.loads(entreprise.contacts_json or "[]") if c.get("nom")]
        for contact in contacts or [{}]:
            feuille.append([_cellule(x) for x in (
                entreprise.nom, entreprise.secteur, contact.get("nom", "non trouvé"),
                contact.get("poste", ""), entreprise.region,
                contact.get("courriel") or "non trouvé", contact.get("source", ""),
                contact.get("source_courriel", ""), entreprise.note)])
    bilan = livre.create_sheet("Bilan")
    bilan.append(["Dossier", _cellule(dossier.nom)])
    bilan.append(["Entreprises traitées", len(entreprises)])
    bilan.append(["Entreprises importées", len(dossier.entreprises)])
    fin = dossier.fin or datetime.utcnow()
    bilan.append(["Durée écoulée (heures)", round((fin - dossier.debut).total_seconds() / 3600, 2)])
    for feuille_courante in livre:
        feuille_courante.freeze_panes = "A2"
        feuille_courante.row_dimensions[1].height = 24
        for cellule in feuille_courante[1]:
            cellule.fill = PatternFill("solid", fgColor="1F3A5F")
            cellule.font = Font(bold=True, color="FFFFFF")
        for col in feuille_courante.columns:
            lettre = col[0].column_letter
            feuille_courante.column_dimensions[lettre].width = min(55, max(18, max(len(str(c.value or "")) for c in col) + 2))
    sortie = io.BytesIO()
    livre.save(sortie)
    return Response(sortie.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="dossier-{dossier.id}.xlsx"'})
