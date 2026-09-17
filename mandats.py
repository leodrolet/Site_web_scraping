"""Validation, import et export des mandats. Aucun accès réseau dans ce module."""
import csv
import io
import json
import re
import unicodedata
import zipfile
from datetime import date
from urllib.parse import urlencode, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ETATS = {"a_faire": "À faire", "en_cours": "En cours", "termine": "Terminé", "a_verifier": "À vérifier"}
VERIFICATIONS = {"a_verifier": "À vérifier", "verifie": "Vérifié", "ecarte": "Écarté"}
COURRIELS = {"non_trouve": "Courriel non trouvé", "a_verifier": "Courriel à vérifier", "trouve": "Courriel trouvé"}
TAILLES = {"inconnue": "Inconnue (cible : 2)", "petite": "Petite (cible : 2 à 3)", "grande": "Grande (cible : 3 à 5)"}
CHAMPS_IMPORT = {
    "nom": ("Organisation", ["entreprise", "nom", "nomentreprise", "nomdelentreprise", "company", "companyname", "organisation", "organization"]),
    "secteur": ("Secteur", ["secteur", "industry", "sector", "industrie"]),
    "pays": ("Pays", ["pays", "country"]),
    "region": ("Région / ville", ["region", "ville", "city", "province", "location"]),
    "site": ("Site web", ["site", "siteweb", "website", "url", "domaine", "domain"]),
    "prix": ("Prix remporté", ["prix", "prixremporte", "award", "awards", "distinction"]),
    "note": ("Notes", ["note", "notes", "commentaire", "comments", "description"]),
    "sources": ("Liens de source", ["source", "sources", "liensource", "sourceurl"]),
    "durabilite": ("Construction durable / carbone zéro", ["durabilite", "durable", "sustainability", "carbonzero", "carbonezero"]),
}
MAX_OCTETS = 2_000_000
MAX_LIGNES = 2000


def normaliser(texte):
    texte = unicodedata.normalize("NFKD", str(texte or "").casefold())
    return "".join(c for c in texte if c.isalnum() and not unicodedata.combining(c))


def nettoyer(texte):
    return " ".join(unicodedata.normalize("NFKC", str(texte or "")).split())


def champ(form, nom, maximum=2000, requis=False):
    valeur = str(form.get(nom, "")).strip()
    if requis and not valeur:
        raise ValueError(f"Le champ « {nom} » est requis.")
    if len(valeur) > maximum:
        raise ValueError(f"Le champ « {nom} » dépasse {maximum} caractères.")
    return valeur


def url(valeur, domaine=False):
    valeur = valeur.strip()
    if not valeur:
        return ""
    if domaine and "://" not in valeur:
        valeur = "https://" + valeur
    try:
        p = urlparse(valeur)
        valide = (p.scheme in ("http", "https") and p.hostname and "." in p.hostname
                  and not p.username and not p.password and not any(c.isspace() for c in valeur))
    except ValueError:
        valide = False
    if not valide or len(valeur) > 2000:
        raise ValueError("Un lien est invalide : utilisez une adresse complète http:// ou https://.")
    return valeur


def liens(valeur):
    return "\n".join(dict.fromkeys(url(v) for v in valeur.splitlines() if v.strip()))


def organisation(form):
    resultat = {k: champ(form, k, 160 if k in ("secteur", "pays", "region") else 4000)
                for k in CHAMPS_IMPORT}
    resultat["nom"] = nettoyer(champ(form, "nom", 255, True))
    resultat["site"] = url(resultat["site"], domaine=True)
    if len(resultat["site"]) > 500:
        raise ValueError("Le site web dépasse 500 caractères.")
    resultat["sources"] = liens(resultat["sources"])
    resultat["taille"] = champ(form, "taille") or "inconnue"
    resultat["statut"] = champ(form, "statut") or "a_faire"
    if resultat["statut"] not in ETATS or resultat["taille"] not in TAILLES:
        raise ValueError("Statut ou taille invalide.")
    resultat["cle"] = normaliser(resultat["nom"])[:255]
    return resultat


def contacts_de(ent):
    """Lecture compatible des anciens contacts, jamais promus en vérifiés."""
    resultat = []
    for i, c in enumerate(json.loads(ent.contacts_json or "[]")):
        if not any(c.get(k) for k in ("nom", "poste", "courriel")):
            continue
        c = dict(c)
        c.setdefault("id", f"ancien-{i}")
        c.setdefault("verification", "a_verifier")
        c.setdefault("retenu", False)
        c.setdefault("courriel_statut", "a_verifier" if c.get("courriel") else "non_trouve")
        for k in ("nom", "poste", "departement", "region", "langue", "courriel", "source",
                  "source_courriel", "date_verification", "note", "provenance", "date_recherche"):
            c.setdefault(k, "")
        resultat.append(c)
    return resultat


def contact(form):
    c = {k: champ(form, k, 2000 if k in ("source", "source_courriel", "note") else 255)
         for k in ("nom", "poste", "departement", "region", "langue", "courriel", "source",
                   "source_courriel", "date_verification", "note")}
    if not c["nom"]:
        raise ValueError("Indiquez le nom du contact. Un lien de profil seul reste une piste.")
    c["source"] = liens(c["source"])
    c["source_courriel"] = url(c["source_courriel"])
    c["verification"] = champ(form, "verification") or "a_verifier"
    c["courriel_statut"] = champ(form, "courriel_statut") or "non_trouve"
    c["retenu"] = form.get("retenu") == "on"
    if c["verification"] not in VERIFICATIONS or c["courriel_statut"] not in COURRIELS:
        raise ValueError("Statut de vérification invalide.")
    if c["courriel"] and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", c["courriel"]):
        raise ValueError("Le format du courriel est invalide.")
    if c["courriel_statut"] == "non_trouve" and c["courriel"]:
        raise ValueError("Un courriel est saisi : choisissez « à vérifier » ou « trouvé ».")
    if c["courriel_statut"] != "non_trouve" and not c["courriel"]:
        raise ValueError("Saisissez le courriel ou choisissez « non trouvé ».")
    if c["courriel_statut"] == "trouve" and not c["source_courriel"]:
        raise ValueError("Un courriel trouvé doit être accompagné du lien de sa source fiable.")
    if c["date_verification"]:
        try:
            jour = date.fromisoformat(c["date_verification"])
        except ValueError as exc:
            raise ValueError("La date de vérification est invalide.") from exc
        if jour > date.today():
            raise ValueError("La date de vérification ne peut pas être dans le futur.")
    if c["verification"] == "verifie" and not (c["poste"] and c["source"] and c["date_verification"]):
        raise ValueError("Pour vérifier un contact, indiquez son poste, les liens confirmant son identité et son poste, et la date de vérification.")
    if c["retenu"] and c["verification"] == "ecarte":
        raise ValueError("Un contact écarté ne peut pas être retenu.")
    return c


def lire_fichier(fichier, ligne_entete=1, feuille=1):
    nom = (fichier.filename or "").lower()
    if not nom.endswith((".csv", ".xlsx")):
        raise ValueError("Format non pris en charge. Utilisez CSV ou Excel .xlsx (pas .xls), ou ajoutez les organisations manuellement.")
    brut = fichier.file.read(MAX_OCTETS + 1)
    if len(brut) > MAX_OCTETS:
        raise ValueError("Le fichier dépasse 2 Mo. Divisez la liste en plusieurs fichiers.")
    if not 1 <= ligne_entete <= 20 or not 1 <= feuille <= 50:
        raise ValueError("La ligne d’en-tête doit être entre 1 et 20 et le numéro de feuille entre 1 et 50.")
    livre = None
    try:
        if nom.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(brut)) as archive:
                if sum(z.file_size for z in archive.infolist()) > 30_000_000:
                    raise ValueError("Le classeur décompressé est trop volumineux. Exportez la feuille en CSV.")
            livre = load_workbook(io.BytesIO(brut), read_only=True, data_only=False)
            if feuille > len(livre.worksheets):
                raise ValueError(f"Ce classeur contient {len(livre.worksheets)} feuille(s).")
            ws = livre.worksheets[feuille - 1]
            if (ws.max_column or 0) > 100:
                raise ValueError("Limite de 100 colonnes. Exportez seulement les colonnes utiles.")
            def valeurs():
                for row in ws.iter_rows():
                    if any(c.data_type == "f" for c in row):
                        raise ValueError("La feuille contient des formules. Collez les valeurs dans un nouveau fichier avant l’import.")
                    yield [str(c.value) if c.value is not None else "" for c in row]
            lignes = valeurs()
        else:
            try:
                contenu = brut.decode("utf-8-sig")
            except UnicodeDecodeError:
                contenu = brut.decode("cp1252")
            try:
                dialecte = csv.Sniffer().sniff(contenu[:8192], delimiters=",;\t")
            except csv.Error:
                dialecte = csv.excel
            lignes = csv.reader(io.StringIO(contenu), dialect=dialecte, strict=True)
        for _ in range(ligne_entete - 1):
            next(lignes, None)
        entetes = next(lignes, [])
        if not entetes or len(entetes) > 100:
            raise ValueError("En-tête vide ou plus de 100 colonnes. Vérifiez le numéro de ligne choisi.")
        entetes = [nettoyer(v) or f"Colonne {i + 1}" for i, v in enumerate(entetes)]
        donnees = []
        for n, row in enumerate(lignes, 1):
            if n > MAX_LIGNES:
                raise ValueError("Limite de 2 000 lignes par fichier. Divisez votre liste.")
            if len(row) > len(entetes):
                raise ValueError(f"Ligne {n + ligne_entete} : plus de valeurs que de colonnes. Vérifiez le séparateur du CSV.")
            row = [str(v).strip() for v in row]
            if any(len(v) > 4000 for v in row):
                raise ValueError(f"Ligne {n + ligne_entete} : une cellule dépasse 4 000 caractères.")
            if any(row):
                donnees.append(row + [""] * (len(entetes) - len(row)))
        if not donnees:
            raise ValueError("Aucune organisation à importer. Vérifiez la feuille et la ligne d’en-tête.")
        return {"entetes": entetes, "lignes": donnees}
    except (ValueError, UnicodeDecodeError) as exc:
        raise ValueError(str(exc)) from exc
    except Exception as exc:
        raise ValueError("Fichier illisible. Réexportez-le en CSV ou .xlsx, ou utilisez la saisie manuelle.") from exc
    finally:
        if livre:
            livre.close()


def proposer_mapping(entetes):
    return {k: next((str(i) for i, h in enumerate(entetes) if normaliser(h) in alias), "")
            for k, (_, alias) in CHAMPS_IMPORT.items()}


def valider_mapping(form, entetes):
    mapping = {k: str(form.get(k, "")) for k in CHAMPS_IMPORT}
    choisis = []
    for valeur in mapping.values():
        if valeur and (not valeur.isdigit() or int(valeur) >= len(entetes)):
            raise ValueError("Correspondance de colonne invalide.")
        if valeur:
            choisis.append(valeur)
    if not mapping["nom"]:
        raise ValueError("Choisissez la colonne contenant le nom de l’organisation.")
    if len(choisis) != len(set(choisis)):
        raise ValueError("Une colonne ne peut pas être associée à deux champs.")
    return mapping


def trouver_doublon(ligne, entreprises):
    for e in entreprises:
        if normaliser(e["nom"]) == normaliser(ligne["nom"]):
            if not e.get("pays") or not ligne.get("pays") or normaliser(e["pays"]) == normaliser(ligne["pays"]):
                return e
    return None


def preparer_import(donnees, mapping, entreprises):
    existantes = [dict(id=e.id, nom=e.nom, pays=e.pays, site=e.site) for e in entreprises]
    resultat = []
    for i, row in enumerate(donnees["lignes"]):
        d = {k: row[int(col)] if col != "" else "" for k, col in mapping.items()}
        d["nom"] = nettoyer(d["nom"])
        avertissements = []
        if not d["nom"]:
            resultat.append(dict(index=i, valeurs=d, etat="invalide", detail="Nom manquant : ligne ignorée."))
            continue
        try:
            d = organisation(d)
        except ValueError as exc:
            resultat.append(dict(index=i, valeurs=d, etat="invalide", detail=str(exc)))
            continue
        doublon = trouver_doublon(d, existantes)
        if doublon:
            etat = "doublon"
            avertissements.append("Même nom et pays compatible : compléter la fiche ou ignorer cette ligne.")
        else:
            etat = "nouvelle"
            if d["site"] and any(e.get("site") and urlparse(e["site"]).hostname == urlparse(d["site"]).hostname for e in existantes):
                avertissements.append("Site partagé avec une autre organisation : vérifier s’il s’agit d’une filiale.")
                d["statut"] = "a_verifier"
            existantes.append(d)
        resultat.append(dict(index=i, valeurs=d, etat=etat, detail=" ".join(avertissements)))
    return resultat


def recherches(ent):
    cible = '"communications" OR "partnerships" OR "partenariats"' if any(
        m in normaliser(ent.secteur) for m in ("universit", "college", "gouvernement", "public", "enseignement")) else '"marketing" OR "sales" OR "business development"'
    nom = f'"{ent.nom}"'
    domaine = urlparse(ent.site).hostname if ent.site else None
    local = f"site:{domaine}" if domaine else nom
    questions = [
        ("Trouver le site officiel", nom + " official website"),
        ("Équipe et direction", local + ' (team OR leadership OR équipe OR direction)'),
        ("Communiqués", local + ' (news OR press OR communiqué)'),
        ("Profils professionnels publics", f'site:linkedin.com/in/ {nom} ({cible})'),
        ("Responsables Canada / région", f'{nom} ({cible}) "{ent.region or "Canada"}"'),
        ("Construction durable / carbone zéro", local + ' (sustainable OR "net zero" OR "carbone zéro")'),
    ]
    return [(label, "https://www.google.com/search?" + urlencode({"q": q})) for label, q in questions]


def statistiques(entreprises):
    contacts = [c for e in entreprises for c in contacts_de(e)]
    traitees = sum(e.statut == "termine" for e in entreprises)
    return dict(total=len(entreprises), traitees=traitees, restantes=len(entreprises) - traitees,
                retenus=sum(bool(c["retenu"]) for c in contacts),
                a_verifier=sum(e.statut == "a_verifier" or any(c["verification"] != "ecarte" and
                    (c["verification"] != "verifie" or c["courriel_statut"] == "a_verifier")
                    for c in contacts_de(e)) for e in entreprises))


def excel(mandat, entreprises):
    livre = Workbook()
    ws = livre.active
    ws.title = "Contacts retenus"
    ws.append(["Organisation", "Secteur", "Pays", "Prix remporté", "Nom", "Poste", "Département",
               "Région du contact", "Langue connue", "Courriel", "Statut courriel", "Sources identité et poste",
               "Source courriel", "Date de vérification", "Vérification", "Note contact", "Provenance", "Date de recherche"])
    def ajouter(feuille, valeurs):
        # Force toutes les données externes en texte, même celles commençant par '='.
        feuille.append(valeurs)
        for cellule in feuille[feuille.max_row]:
            if isinstance(cellule.value, str):
                cellule.data_type = "s"
                cellule.alignment = Alignment(vertical="top", wrap_text=True)
    for e in entreprises:
        for c in contacts_de(e):
            if c["retenu"]:
                ajouter(ws, [e.nom, e.secteur, e.pays, e.prix, c["nom"], c["poste"], c["departement"],
                             c["region"], c["langue"] or "Inconnue", c["courriel"] or "non trouvé",
                             COURRIELS[c["courriel_statut"]], c["source"], c["source_courriel"],
                             c["date_verification"], VERIFICATIONS[c["verification"]], c["note"],
                             c["provenance"], c["date_recherche"]])
    orgs = livre.create_sheet("Organisations")
    orgs.append(["Organisation", "Statut", "Secteur", "Pays", "Région", "Site", "Prix", "Durabilité",
                 "Sources", "Notes", "Contacts retenus", "Dernière sauvegarde UTC", "Résultat"])
    for e in entreprises:
        n = sum(bool(c["retenu"]) for c in contacts_de(e))
        ajouter(orgs, [e.nom, ETATS[e.statut], e.secteur, e.pays, e.region, e.site, e.prix,
                       e.durabilite, e.sources, e.note, n,
                       e.mise_a_jour.isoformat() if e.mise_a_jour else "",
                       "Aucun contact pertinent retenu" if e.statut == "termine" and not n else ""])
    bilan = livre.create_sheet("Bilan")
    bilan.append(["Indicateur", "Valeur"])
    ajouter(bilan, ["Mandat", mandat.nom])
    for cle, valeur in statistiques(entreprises).items():
        ajouter(bilan, [{"total": "Organisations au total", "traitees": "Organisations traitées",
                         "restantes": "Organisations restantes", "retenus": "Contacts retenus",
                         "a_verifier": "Fiches à vérifier"}[cle], valeur])
    ajouter(bilan, ["Temps de travail saisi (minutes)", mandat.temps_minutes])
    ajouter(bilan, ["Budget du mandat (minutes)", mandat.budget_minutes])
    for feuille in livre:
        feuille.freeze_panes = "A2"
        feuille.auto_filter.ref = feuille.dimensions
        for c in feuille[1]:
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.font = Font(bold=True, color="FFFFFF")
        for col in feuille.columns:
            feuille.column_dimensions[col[0].column_letter].width = min(55, max(20, max(len(str(c.value or "")) for c in col) + 2))
    sortie = io.BytesIO()
    livre.save(sortie)
    return sortie.getvalue()
