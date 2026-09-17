"""Imports guidés sans réseau ni logiciel Office : OOXML et Excel, données non fiables."""
import io
import re
import zipfile
from xml.dom import minidom
from openpyxl import load_workbook

MAX_BYTES = 2_000_000
MAX_ITEMS = 500
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
EMAIL = re.compile(r'[^\s<>;,]+@[^\s<>;,]+\.[^\s<>;,]+')
ROLE = re.compile(r'\b(director|manager|president|sales|marketing|mktg|communications?|partnerships?|partenariats?|ventes|directeur|directrice|responsable|executive|specialist|management|mgr)\b', re.I)
LABEL = re.compile(r'^(emails?|website|phone numbers|main contact details|additional contact details|contact details|location|local|\([A-Z]+\)|A\+?|TRY;?)$', re.I)


def archive(brut, partie):
    if len(brut) > MAX_BYTES:
        raise ValueError('Le fichier dépasse 2 Mo. Divisez le document avant de réessayer.')
    try:
        z = zipfile.ZipFile(io.BytesIO(brut))
        infos = z.infolist()
        if len(infos) > 2000 or sum(x.file_size for x in infos) > 30_000_000:
            raise ValueError('Document décompressé trop volumineux (30 Mo maximum).')
        if len({x.filename for x in infos}) != len(infos) or any(x.flag_bits & 1 for x in infos):
            raise ValueError('Archive ambiguë ou chiffrée non prise en charge.')
        if partie not in z.namelist() or any('vbaproject' in x.filename.lower() or x.filename.startswith('word/embeddings/') for x in infos):
            raise ValueError('Format incorrect, macro ou objet incorporé non pris en charge. Enregistrez une copie simple sans macros.')
        for x in infos:
            if x.filename.endswith(('.xml', '.rels')):
                contenu = z.read(x)
                if b'\x00' in contenu or b'<!DOCTYPE' in contenu.upper() or b'<!ENTITY' in contenu.upper():
                    raise ValueError('Déclarations XML externes interdites.')
        return z
    except (zipfile.BadZipFile, RuntimeError, OSError) as exc:
        raise ValueError('Fichier illisible ou mal formé. Réenregistrez une copie avec Word ou Excel.') from exc


def xml(brut):
    try:
        if b'\x00' in brut or b'<!DOCTYPE' in brut.upper() or b'<!ENTITY' in brut.upper():
            raise ValueError('Déclarations XML interdites.')
        return minidom.parseString(brut)
    except Exception as exc:
        raise ValueError('Le document XML est mal formé.') from exc


def nodes(n, tag):
    return n.getElementsByTagNameNS(W, tag)


def attr(n, key='val'):
    return n.getAttributeNS(W, key)


def texte(n):
    return ''.join(''.join(c.data for c in t.childNodes if c.nodeType == c.TEXT_NODE) for t in nodes(n, 't'))


def est_rouge(val):
    if not re.fullmatch('[0-9a-fA-F]{6}', val):
        return False
    r, g, b = (int(val[i:i+2], 16) for i in (0, 2, 4))
    return r >= 120 and r > g * 1.5 and r > b * 1.5


def lire_word(brut):
    with archive(brut, 'word/document.xml') as z:
        d = xml(z.read('word/document.xml'))
        if d.documentElement.namespaceURI != W or d.documentElement.localName != 'document' or not nodes(d, 'body'):
            raise ValueError('Le fichier ne contient pas un document Word OOXML valide.')
        styles = {}
        if 'word/styles.xml' in z.namelist():
            for s in nodes(xml(z.read('word/styles.xml')), 'style'):
                styles[attr(s, 'styleId')] = s
        theme = {}
        if 'word/theme/theme1.xml' in z.namelist():
            td = xml(z.read('word/theme/theme1.xml'))
            for scheme in td.getElementsByTagNameNS('*', 'clrScheme'):
                for c in scheme.childNodes:
                    if c.nodeType == c.ELEMENT_NODE:
                        children = [x for x in c.childNodes if x.nodeType == x.ELEMENT_NODE]
                        if children:
                            theme[c.localName] = children[0].getAttribute('val') or children[0].getAttribute('lastClr')

        def couleur(n):
            cs = nodes(n, 'color')
            if not cs:
                return None
            c = cs[0]
            return theme.get(attr(c, 'themeColor'), attr(c))

        def style_color(ident, vus=None):
            vus = set() if vus is None else vus
            if ident in vus or ident not in styles:
                return None
            vus.add(ident)
            s = styles[ident]
            parent = nodes(s, 'basedOn')
            return couleur(s) or (style_color(attr(parent[0]), vus) if parent else None)

        ps = list(nodes(d, 'p'))
        if len(ps) > 5000:
            raise ValueError('Document trop long : 5 000 paragraphes maximum.')
        items = []
        for index, p in enumerate(ps):
            pstyles = nodes(p, 'pStyle')
            inherited = style_color(attr(pstyles[0])) if pstyles else style_color('Normal')
            segments, actuel = [], ''
            for r in nodes(p, 'r'):
                rs = nodes(r, 'rStyle')
                col = couleur(r) or (style_color(attr(rs[0])) if rs else None) or inherited or ''
                if est_rouge(col):
                    actuel += texte(r)
                elif actuel:
                    segments.append(actuel.strip()); actuel = ''
            if actuel:
                segments.append(actuel.strip())
            segments = [s for s in segments if s]
            if not segments:
                continue
            red = ' / '.join(segments)
            if len(texte(p)) > 12000:
                raise ValueError('Un paragraphe dépasse 12 000 caractères. Divisez-le avant import.')
            suspect = bool(EMAIL.search(red) or ROLE.search(red) or re.match(r'^(Ms|Mr|Mme|M\.)\s', red))
            items.append(dict(ancre=index, rouge=red, contexte=texte(p), avant=texte(ps[index-1]) if index else '',
                apres=texte(ps[index+1]) if index+1 < len(ps) else '', nom='' if suspect or len(segments) != 1 else red,
                contacts=[], note='', etat='a_confirmer', organisation_id=None,
                avertissement='Ce passage peut déjà désigner une personne ou plusieurs éléments. Choisissez une organisation ou ignorez-le.' if suspect or len(segments) != 1 else 'Nom proposé à confirmer : le rouge ne prouve pas une demande non traitée.'))
        if len(items) > MAX_ITEMS:
            raise ValueError('Plus de 500 passages rouges. Divisez le document.')
        return items


def contact_colonne(vals):
    """Ne devine ni société depuis un domaine, ni langue, ni adresse absente."""
    useful = [(r, v) for r, v in vals if v and not LABEL.fullmatch(v)]
    email = next((EMAIL.search(v).group().rstrip('.') for _, v in useful if EMAIL.search(v)), '')
    job = next((v for _, v in useful if ROLE.search(v)), '')
    nom = ''
    for _, v in useful:
        if EMAIL.search(v) or v.startswith(('http', 'www.')) or re.match(r'^[\d(+]', v) or LABEL.fullmatch(v):
            continue
        if v == job:
            if ',' in v and re.match(r'^(Ms |Mr |Mme )', v):
                nom, job = v.split(',', 1)
            continue
        if len(v.split()) in range(2, 6) and len(v) < 100 and not re.search(r'\b(no |could |try|available|info|furniture)\b', v, re.I):
            nom = v
            break
    region = ''
    if email:
        emailrow = next(r for r, v in useful if email in v)
        apres = [(r,v) for r,v in useful if r > emailrow and v != job and not ROLE.search(v)]
        if apres and len(apres) == 1 and not re.match(r'^[\d(+]', apres[0][1]) and not re.search(r'\b(try|could|available|email|info)\b', apres[0][1], re.I) and len(apres[0][1]) < 160:
            region = apres[0][1]
    source = next((v for _,v in vals if v.startswith(('https://','http://'))), '')
    return dict(nom=nom, poste=job.strip(), courriel=email, region=region, source=source, note='\n'.join(v for _, v in vals)[:2000])


def lire_blocs(brut, feuille=1):
    with archive(brut, 'xl/workbook.xml'):
        pass
    try:
        wb = load_workbook(io.BytesIO(brut), read_only=True, data_only=False)
        try:
            if not 1 <= feuille <= len(wb.worksheets):
                raise ValueError('Numéro de feuille absent du classeur.')
            ws = wb.worksheets[feuille-1]
            if ws.max_row > 2000 or ws.max_column > 100:
                raise ValueError('La feuille dépasse 2 000 lignes ou 100 colonnes.')
            rows = []
            for row in ws.iter_rows():
                cells = []
                for c in row:
                    if c.data_type == 'f':
                        raise ValueError('La feuille contient des formules. Copiez-collez les valeurs dans un nouveau classeur.')
                    if c.value is not None:
                        v = str(c.value).strip()
                        if len(v) > 4000:
                            raise ValueError('Une cellule dépasse 4 000 caractères.')
                        if v:
                            cells.append((c.column, c.row, v))
                rows.append(cells)
            bands, band = [], []
            for row in rows + [[]]:
                if row:
                    band.extend(row)
                elif band:
                    bands.append(band); band = []
            result = []
            for band in bands:
                cols = sorted({c for c, _, _ in band})
                firstrow = min(r for _, r, _ in band)
                first = [(c, v) for c, r, v in band if r == firstrow]
                left = [(r, v) for c, r, v in band if c == cols[0]]
                # A heading followed by a person then a title/email is a proposal only.
                candidate = left[0][1]
                has_heading = (len(left) >= 3 and not ROLE.search(candidate) and not EMAIL.search(candidate)
                    and not candidate.startswith(('http', 'www.')) and not LABEL.fullmatch(candidate)
                    and not ROLE.search(left[1][1]) and (ROLE.search(left[2][1]) or EMAIL.search(left[2][1])))
                nom = candidate if has_heading else ''
                contacts = []
                for col in cols:
                    vals = [(r, v) for c, r, v in band if c == col and not (has_heading and c == cols[0] and r == firstrow)]
                    co = contact_colonne(vals)
                    if co['nom'] and (co['poste'] or co['courriel']):
                        contacts.append(co)
                context = '\n'.join(f'{ws.cell(r,c).coordinate} : {v}' for c,r,v in sorted(band,key=lambda x:(x[1],x[0])))
                result.append(dict(nom=nom, contacts=contacts[:12], contexte=context, rouge='', avant='', apres='',
                    ancre=f'{ws.title} · lignes {min(r for _,r,_ in band)}–{max(r for _,r,_ in band)}',
                    note='', etat='a_confirmer', organisation_id=None,
                    avertissement='Vérifiez entreprise, personnes, fonctions et adresses. Les mentions « try / could try » restent dans les notes; aucune adresse importée n’est confirmée.' if nom else 'Organisation ambiguë ou absente : saisissez-la ou choisissez une fiche existante. Le domaine d’un courriel ne suffit pas à identifier une entreprise.'))
            if len(result) > MAX_ITEMS:
                raise ValueError('Plus de 500 blocs. Divisez la feuille.')
            return result
        finally:
            wb.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Classeur mal formé ou illisible. Réenregistrez-le en .xlsx.') from exc


def lignes_reponse(contacts):
    lignes = []
    for c in contacts:
        if not c.get('retenu') or c.get('verification') == 'ecarte':
            continue
        parts = [c['nom'], c.get('poste'), c.get('courriel') or 'courriel non trouvé', c.get('region')]
        if c.get('verification') != 'verifie':
            parts.append('contact à vérifier')
        if c.get('courriel') and c.get('courriel_statut') != 'trouve':
            parts.append('courriel à vérifier')
        lignes.append(' — '.join(v for v in parts if v))
    return lignes


def completer_word(brut, insertions):
    """Conserve les autres parties OOXML à l'identique; ajoute du texte rouge dans le paragraphe lié."""
    with archive(brut, 'word/document.xml') as z:
        d = xml(z.read('word/document.xml'))
        d.documentElement.setAttribute('xmlns:w', W)
        ps = list(nodes(d, 'p'))
        for index, lignes in insertions.items():
            if not 0 <= int(index) < len(ps):
                raise ValueError('Emplacement Word introuvable.')
            p = ps[int(index)]
            for ligne in lignes:
                if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ligne):
                    raise ValueError('Un contact contient un caractère de contrôle incompatible avec Word. Corrigez sa fiche.')
                r = d.createElementNS(W, 'w:r')
                rp = d.createElementNS(W, 'w:rPr')
                color = d.createElementNS(W, 'w:color'); color.setAttributeNS(W, 'w:val', 'FF0000')
                rp.appendChild(color); r.appendChild(rp)
                r.appendChild(d.createElementNS(W, 'w:br'))
                t = d.createElementNS(W, 'w:t'); t.appendChild(d.createTextNode(ligne)); r.appendChild(t)
                p.appendChild(r)
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
            for entry in z.infolist():
                dst.writestr(entry, d.toxml(encoding='UTF-8') if entry.filename == 'word/document.xml' else z.read(entry))
        return out.getvalue()
